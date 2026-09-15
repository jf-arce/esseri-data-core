"""RF-37: exportación del historial de asistencias.

`GET /academico/asistencias/exportar` es un reporte institucional (cruza todos los alumnos de
una división, o de toda la institución sin `division_id`) -- a diferencia de
`listar_asistencias`/`calcular_resumen_asistencia` (RF-05/06), acotadas a una inscripción
puntual. El único rol pensado para esto es `dirección` (`Académico: leer, exportar`, sin
`actualizar` de ningún tipo), que antes de RF-37 quedaba bloqueada por el scoping por división
de un docente -- ver el comentario en `_tiene_acceso_estructural` (`academico/service.py`).
"""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from src.familias_alumnos.models import Alumno
from src.inscripciones.models import Asistencia, Inscripcion
from src.models import Persona
from tests.inscripciones.factories import (
    crear_division_destino,
    crear_escenario,
    crear_inscripcion_previa,
)


def _cargar_asistencias_de_prueba(db_session, inscripcion_id):
    db_session.add_all(
        [
            Asistencia(inscripcion_id=inscripcion_id, fecha=date(2027, 3, 1), tipo="presente"),
            Asistencia(inscripcion_id=inscripcion_id, fecha=date(2027, 3, 2), tipo="tardanza"),
            Asistencia(
                inscripcion_id=inscripcion_id, fecha=date(2027, 3, 3), tipo="ausente_justificado"
            ),
        ]
    )
    db_session.commit()


def test_docente_no_puede_exportar_asistencias(client_docente):
    """El docente nunca tiene `academico.exportar` -- ni siquiera con una división asignada."""
    respuesta = client_docente.get(
        "/academico/asistencias/exportar?fecha_desde=2027-03-01&fecha_hasta=2027-03-31"
    )

    assert respuesta.status_code == 403


def test_secretaria_no_puede_exportar_asistencias(client_secretaria):
    """Secretaría tiene `actualizar` estructural amplio, pero no `exportar` -- son permisos
    independientes, uno no sustituye al otro."""
    respuesta = client_secretaria.get(
        "/academico/asistencias/exportar?fecha_desde=2027-03-01&fecha_hasta=2027-03-31"
    )

    assert respuesta.status_code == 403


def test_direccion_puede_exportar_todas_las_divisiones(client_direccion, db_session):
    escenario = crear_escenario(db_session)
    inscripcion = crear_inscripcion_previa(db_session, escenario, estado="activa")
    _cargar_asistencias_de_prueba(db_session, inscripcion.id)

    respuesta = client_direccion.get(
        "/academico/asistencias/exportar?fecha_desde=2027-03-01&fecha_hasta=2027-03-31"
    )

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "text/csv; charset=utf-8"
    assert (
        'filename="asistencias_2027-03-01_2027-03-31.csv"'
        in (respuesta.headers["content-disposition"])
    )
    cuerpo = respuesta.content.decode("utf-8-sig")
    assert "Fecha,Alumno,Legajo,División,Estado" in cuerpo
    assert "Cabral, Tiziano" in cuerpo
    assert "Presente" in cuerpo
    assert "Ausente justificado" in cuerpo


def test_direccion_puede_acotar_la_exportacion_a_una_division(client_direccion, db_session):
    escenario = crear_escenario(db_session)
    inscripcion = crear_inscripcion_previa(db_session, escenario, estado="activa")
    _cargar_asistencias_de_prueba(db_session, inscripcion.id)

    # Segundo alumno en OTRA división del mismo año, sin volver a llamar crear_escenario (que
    # crea un Usuario con email fijo -- llamarlo dos veces rompe por UNIQUE).
    otra_division_id = crear_division_destino(db_session, escenario)
    otra_persona = Persona(nombre="Otro", apellido="Alumno", dni="50333444")
    db_session.add(otra_persona)
    db_session.flush()
    otro_alumno = Alumno(numero_legajo="A-2027-002", estado="activo", persona_id=otra_persona.id)
    db_session.add(otro_alumno)
    db_session.flush()
    otra_inscripcion = Inscripcion(
        ciclo_lectivo="2027",
        fecha_inscripcion=date(2026, 8, 27),
        tipo="nueva",
        estado="activa",
        alumno_id=otro_alumno.id,
        division_id=otra_division_id,
        solicitud_inscripcion_id=escenario["solicitud_id"],
    )
    db_session.add(otra_inscripcion)
    db_session.commit()
    _cargar_asistencias_de_prueba(db_session, otra_inscripcion.id)

    respuesta = client_direccion.get(
        "/academico/asistencias/exportar"
        f"?fecha_desde=2027-03-01&fecha_hasta=2027-03-31&division_id={escenario['division_id']}"
    )

    assert respuesta.status_code == 200
    cuerpo = respuesta.content.decode("utf-8-sig")
    assert "Cabral, Tiziano" in cuerpo
    assert "Otro" not in cuerpo


def test_exportar_en_excel_genera_un_libro_valido(client_direccion, db_session):
    escenario = crear_escenario(db_session)
    inscripcion = crear_inscripcion_previa(db_session, escenario, estado="activa")
    _cargar_asistencias_de_prueba(db_session, inscripcion.id)

    respuesta = client_direccion.get(
        "/academico/asistencias/exportar?fecha_desde=2027-03-01&fecha_hasta=2027-03-31&formato=xlsx"
    )

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="asistencias_2027-03-01_2027-03-31.xlsx"'
        in (respuesta.headers["content-disposition"])
    )

    libro = load_workbook(BytesIO(respuesta.content))
    hoja = libro.active
    encabezado = [celda.value for celda in next(hoja.iter_rows(max_row=1))]
    assert encabezado == ["Fecha", "Alumno", "Legajo", "División", "Estado"]
    # 3 filas de datos + encabezado
    assert hoja.max_row == 4


def test_exportar_en_pdf_genera_un_documento_valido(client_direccion, db_session):
    escenario = crear_escenario(db_session)
    inscripcion = crear_inscripcion_previa(db_session, escenario, estado="activa")
    _cargar_asistencias_de_prueba(db_session, inscripcion.id)

    respuesta = client_direccion.get(
        "/academico/asistencias/exportar?fecha_desde=2027-03-01&fecha_hasta=2027-03-31&formato=pdf"
    )

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert respuesta.content.startswith(b"%PDF")
